#!/usr/bin/env python3
"""Summarize CAMBC run-batch logs against the generated match manifest.

Works with UTF-16 (common on Windows PowerShell/native output) and UTF-8 logs.
Usage: python scripts/summarize_tournament_log.py [--manifest FILE] [--log FILE]
If omitted, chooses the newest root matches-*.csv and the newest compatible log.
"""
import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def read_text_auto(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16', errors='replace')
    # Some Windows logs are UTF-16 LE without a BOM; lots of NULs are a clue.
    if raw and raw.count(b'\x00') / len(raw) > 0.10:
        return raw.decode('utf-16-le', errors='replace')
    return raw.decode('utf-8-sig', errors='replace')


def read_csv(path: Path):
    with path.open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def timestamp_key(path: Path):
    match = re.search(r'(\d{8}-\d{6})', path.name)
    return match.group(1) if match else ''


def choose_files(manifest_arg, log_arg):
    manifest = Path(manifest_arg).resolve() if manifest_arg else None
    if manifest is None:
        choices = list(ROOT.glob('matches-*.csv'))
        if not choices:
            raise SystemExit('No root matches-*.csv manifest found. Run Run-Tournament.ps1 first.')
        manifest = max(choices, key=lambda p: (timestamp_key(p), p.stat().st_mtime))
    if not manifest.exists():
        raise SystemExit(f'Manifest not found: {manifest}')
    log = Path(log_arg).resolve() if log_arg else None
    if log is None:
        # Prefer a log whose mode/coverage and timestamp correspond to the manifest.
        stem = manifest.stem
        logs = list((ROOT / 'reports' / 'logs').glob('*.log'))
        suffix = stem.removeprefix('matches-')
        exact = [p for p in logs if suffix in p.name]
        if exact:
            log = max(exact, key=lambda p: (timestamp_key(p), p.stat().st_mtime))
        else:
            prefix = suffix.rsplit('-', 1)[0]
            similar = [p for p in logs if prefix in p.name]
            if similar:
                log = max(similar, key=lambda p: (timestamp_key(p), p.stat().st_mtime))
    if log is None or not log.exists():
        raise SystemExit('Could not find a matching tournament log. Pass --log reports/logs/<logfile>.log')
    return manifest, log


def save_csv(path, rows):
    if not rows:
        return
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--manifest', help='CSV manifest created by Run-Tournament.ps1')
    ap.add_argument('--log', help='run-batch log created by Run-Tournament.ps1')
    ap.add_argument('--candidate', default='candidate-bot', help='candidate name for head-to-head/by-map reports')
    args = ap.parse_args()
    manifest, log = choose_files(args.manifest, args.log)
    planned = read_csv(manifest)
    if not planned:
        raise SystemExit(f'Manifest has no rows: {manifest}')

    text = read_text_auto(log)
    text = re.sub(r'\x1b\[[0-?]*[ -/]*[@-~]', '', text)
    bots = sorted({r.get('bot_a', '') for r in planned} | {r.get('bot_b', '') for r in planned}, key=len, reverse=True)
    bots = [b for b in bots if b]
    if not bots:
        raise SystemExit('Manifest is missing bot_a/bot_b columns.')
    bot_re = '(?:' + '|'.join(re.escape(b) for b in bots) + ')'
    pattern = re.compile(
        r'TEAM\s+(?P<seat>[AB])\s+WIN\s+\((?P<winner>[^)]+)\)\s+'
        r'(?P<a>' + bot_re + r')\s+vs\s+(?P<b>' + bot_re + r')\s+'
        r'(?P<map>[A-Za-z0-9_.-]+)\s+Round\s+(?P<round>\d+)\s*/\s*(?P<maxround>\d+)',
        re.IGNORECASE,
    )
    parsed = list(pattern.finditer(text))
    if not parsed:
        sample = next((line for line in text.splitlines() if 'WIN' in line.upper()), '(no line containing WIN found)')
        raise SystemExit(f'No winner records parsed from {log}. First possible WIN line: {sample[:300]!r}')

    index = defaultdict(list)
    for row in planned:
        key = (row['bot_a'], row['bot_b'], Path(row['map']).stem)
        index[key].append(row)
    used = Counter()
    results, unmatched = [], []
    for m in parsed:
        d = m.groupdict()
        key = (d['a'], d['b'], d['map'])
        available = [r for r in index.get(key, []) if used[r['id']] == 0]
        if not available:
            unmatched.append(f"No manifest row for {d['a']} vs {d['b']} on {d['map']}")
            continue
        row = available[0]
        winner = d['winner'].strip()
        if winner not in (row['bot_a'], row['bot_b']):
            unmatched.append(f'Unexpected winner {winner} for {key}')
            continue
        used[row['id']] += 1
        results.append({
            'match_id': row['id'], 'map': d['map'], 'bot_a': row['bot_a'], 'bot_b': row['bot_b'],
            'winner': winner, 'loser': row['bot_b'] if winner == row['bot_a'] else row['bot_a'],
            'winner_seat': 'A' if winner == row['bot_a'] else 'B',
            'round': int(d['round']), 'max_round': int(d['maxround']),
        })

    all_bots = sorted({r['bot_a'] for r in planned} | {r['bot_b'] for r in planned})
    stats = {b: Counter() for b in all_bots}
    for r in results:
        stats[r['winner']]['wins'] += 1
        stats[r['loser']]['losses'] += 1
        stats[r['bot_a']]['matches'] += 1
        stats[r['bot_b']]['matches'] += 1
        stats[r['winner']]['wins_as_' + r['winner_seat']] += 1
    standings = []
    for b, s in stats.items():
        n = s['matches']
        standings.append({'bot': b, 'matches': n, 'wins': s['wins'], 'losses': s['losses'],
                          'win_pct': round(100*s['wins']/n, 2) if n else 0,
                          'wins_as_A': s['wins_as_A'], 'wins_as_B': s['wins_as_B']})
    standings.sort(key=lambda x: (-x['win_pct'], -x['wins'], x['bot']))

    out = ROOT / 'reports'
    out.mkdir(exist_ok=True)
    result_path = out / 'tournament-match-results.csv'
    standings_path = out / 'tournament-standings.csv'
    save_csv(result_path, results)
    save_csv(standings_path, standings)
    candidate = args.candidate
    h2h = []
    for opponent in all_bots:
        if opponent == candidate:
            continue
        subset = [r for r in results if {r['bot_a'], r['bot_b']} == {candidate, opponent}]
        wins = sum(r['winner'] == candidate for r in subset)
        losses = sum(r['loser'] == candidate for r in subset)
        h2h.append({'candidate': candidate, 'opponent': opponent, 'matches': len(subset), 'wins': wins,
                    'losses': losses, 'win_pct': round(100*wins/len(subset), 2) if subset else 0})
    h2h_path = out / 'candidate-head-to-head.csv'
    save_csv(h2h_path, h2h)
    by_map = defaultdict(list)
    for r in results:
        if candidate in (r['bot_a'], r['bot_b']):
            by_map[r['map']].append(r)
    map_rows = []
    for name, subset in sorted(by_map.items()):
        wins = sum(r['winner'] == candidate for r in subset)
        losses = sum(r['loser'] == candidate for r in subset)
        map_rows.append({'map': name, 'matches': len(subset), 'wins': wins, 'losses': losses,
                         'win_pct': round(100*wins/len(subset), 2) if subset else 0})
    map_path = out / 'candidate-by-map.csv'
    save_csv(map_path, map_rows)

    # Convenient Excel workbook; CSV outputs remain the canonical simple format.
    xlsx_path = out / 'tournament-report.xlsx'
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        for title, rows in [('Standings', standings), ('Head-to-head', h2h), ('By-map', map_rows), ('Match results', results)]:
            ws = wb.active if title == 'Standings' else wb.create_sheet()
            ws.title = title
            if rows:
                ws.append(list(rows[0].keys()))
                for row in rows: ws.append([row[k] for k in rows[0].keys()])
                ws.freeze_panes = 'A2'
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='244062')
                    cell.alignment = Alignment(wrap_text=True)
                for col in ws.columns:
                    width = min(34, max(10, max(len(str(c.value or '')) for c in col) + 2))
                    ws.column_dimensions[get_column_letter(col[0].column)].width = width
        wb.save(xlsx_path)
    except ImportError:
        xlsx_path = None

    planned_ids = {r['id'] for r in planned}
    matched_ids = {r['match_id'] for r in results}
    print(f'Manifest: {manifest.relative_to(ROOT) if manifest.is_relative_to(ROOT) else manifest}')
    print(f'Log: {log.relative_to(ROOT) if log.is_relative_to(ROOT) else log}')
    print('\n=== STANDINGS ===')
    for s in standings:
        print(f"{s['bot']:20} {s['wins']:4} W - {s['losses']:4} L | {s['win_pct']:6.2f}% | {s['matches']} matches")
    print('\n=== CANDIDATE VS OPPONENTS ===')
    for r in h2h:
        print(f"{r['opponent']:20} {r['wins']:4} W - {r['losses']:4} L | {r['win_pct']:6.2f}% | {r['matches']} matches")
    print('\n=== VALIDATION ===')
    print(f'Planned matches: {len(planned)} | winner lines: {len(parsed)} | matched outcomes: {len(results)}')
    print(f'Unmatched winner lines: {len(unmatched)} | planned without outcome: {len(planned_ids - matched_ids)}')
    print(f'Reports: {result_path.name}, {standings_path.name}, {h2h_path.name}, {map_path.name}')
    if xlsx_path: print(f'Excel workbook: {xlsx_path.relative_to(ROOT) if xlsx_path.is_relative_to(ROOT) else xlsx_path}')
    if unmatched:
        print('\nUnmatched examples:')
        for item in unmatched[:10]: print(' - ' + item)
    if len(results) != len(planned) or unmatched:
        print('\nWARNING: The report is partial. Do not treat missing matches as losses.')
        return 1
    return 0

if __name__ == '__main__':
    sys.exit(main())
